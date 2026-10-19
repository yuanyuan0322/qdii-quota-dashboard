# -*- coding: utf-8 -*-
"""
最终输出：将渠道信息追加到Excel中
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import json
import openpyxl
import os
import ssl
import urllib.request
import re

XL_IN = r'D:\opencode\project\oracle_xrisk\20261013_QDII额度\QDII额度.xlsx'
XL_OUT = r'D:\opencode\project\oracle_xrisk\20261013_QDII额度\QDII额度_渠道标注.xlsx'
RESULT_FILE = r'D:\opencode\project\oracle_xrisk\20261013_QDII额度\_results.json'
TTJJ_DIR = r'D:\opencode\project\oracle_xrisk\20261013_QDII额度\_ttjj_cache'
os.makedirs(TTJJ_DIR, exist_ok=True)

CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE
HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
}


def fetch_ttjj(code):
    code_short = code.replace('.OF', '')
    fn = os.path.join(TTJJ_DIR, f'{code_short}.html')
    if os.path.exists(fn):
        with open(fn, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    url = f'https://fundf10.eastmoney.com/jjjz_{code_short}.html'
    try:
        req = urllib.request.Request(url, headers=HEADERS)
        data = urllib.request.urlopen(req, timeout=15, context=CTX).read()
        with open(fn, 'wb') as f:
            f.write(data)
        return data.decode('utf-8', errors='ignore')
    except Exception:
        return None


def get_ttjj_limit_wan(html):
    """获取天天基金网显示的限额（单位：万元）"""
    if not html:
        return None
    # 匹配 "单日累计购买上限XXXX.XX万元" or "日累计申购限额 XXXX.XX万元"
    m = re.search(r'(?:单日累计购买上限|日累计申购限额)\s*(\d[\d,\.]*\d)\s*万元', html)
    if m:
        try:
            return float(m.group(1).replace(',', ''))
        except Exception:
            return None
    return None


def infer_channel(excel_limit, ttjj_limit_wan):
    """根据Excel和天天基金网限额的关系推断渠道"""
    if ttjj_limit_wan is None:
        # 没法确认，默认全部渠道
        return '全部渠道（含天天基金、支付宝、直销）'
    excel_wan = excel_limit / 10000.0
    if abs(ttjj_limit_wan - excel_wan) < 0.5:
        # 一致 → 全部渠道
        return '全部渠道（含天天基金、支付宝、直销）'
    if ttjj_limit_wan < excel_wan:
        # ttjj 更严 → Excel值是直销或更宽松渠道的限额
        return f'直销为主（天天基金限 {ttjj_limit_wan:g} 万元）'
    # ttjj 更高 → 不太
    return '需人工核实'


def expand_channel(raw):
    """将原始渠道分类展开为更详细的描述"""
    if raw == '全部渠道':
        return '全部渠道（含天天基金、支付宝、直销）'
    if raw == '直销':
        return '直销（基金公司官网/APP）'
    if raw == '代销':
        return '代销（含天天基金、支付宝等）'
    if raw == '直销+代销':
        return '全部渠道（直销+天天基金+支付宝等代销）'
    if raw.startswith('直销+代销（含'):
        special = raw.replace('直销+代销（含', '').replace('）', '')
        return f'全部渠道（直销+天天基金+支付宝+{special}）'
    if raw.startswith('代销（含'):
        special = raw.replace('代销（含', '').replace('）', '')
        return f'代销（天天基金+支付宝+{special}，不含直销）'
    if raw == '未找到公告':
        return None  # 需要额外判断
    return raw


if __name__ == '__main__':
    # 加载结果
    with open(RESULT_FILE, 'r', encoding='utf-8') as f:
        results = json.load(f)
    res_by_code = {r['code']: r for r in results}

    # 读取原Excel
    wb = openpyxl.load_workbook(XL_IN, data_only=True)
    ws = wb.active

    # 添加新列标题
    ws.cell(row=1, column=6).value = '适用渠道（基于公告）'

    # 遍历原数据
    rows_data = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        code = row[0]
        name = row[1]
        limit = row[4]
        if limit is None:
            # 开放申购，不需要标记
            ws.cell(row=i, column=6).value = '正常开放'
            continue
        # 查找渠道
        if code in res_by_code:
            r = res_by_code[code]
            ch = r['channel']
            if ch == '未找到公告':
                # 使用天天基金网页兜底判断
                html = fetch_ttjj(code)
                ttjj_wan = get_ttjj_limit_wan(html)
                ch_out = infer_channel(limit, ttjj_wan)
                ws.cell(row=i, column=6).value = ch_out + '（公告暂未确认）'
            else:
                ch_out = expand_channel(ch)
                ws.cell(row=i, column=6).value = ch_out or ch
        else:
            ws.cell(row=i, column=6).value = '未处理'

    wb.save(XL_OUT)
    print(f'保存到: {XL_OUT}')

    # 显示统计
    from collections import Counter
    chs = Counter()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        if row[5]:
            chs[row[5]] += 1
    print('\n渠道分布:')
    for k, v in chs.most_common():
        print(f'  {k}: {v}')
