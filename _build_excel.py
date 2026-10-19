# -*- coding: utf-8 -*-
"""把 _quota_raw.json 整理成 Excel：天天基金「海外额度宝」额度排行"""
import sys, os, json, re, time
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "_quota_raw.json")
DST = os.path.join(HERE, "天天基金_海外额度宝_额度排行.xlsx")

# ---------- 区域推断（按基金名称关键词） ----------
REGION_RULES = [
    (["纳斯达克", "标普500", "道琼斯", "美国", "美股", "纳指", "美元债"], "美国"),
    (["恒生", "港股", "香港", "中概", "H股", "中国互联", "恒指", "沪港深", "大中华", "港美", "中国"], "大中华地区"),
    (["日本", "日经", "东京"], "日本"),
    (["德国", "法国", "英国", "富时100", "欧洲", "欧元"], "欧洲"),
    (["印度"], "印度"),
    (["越南"], "越南"),
    (["沙特", "中东"], "中东"),
    (["新兴市场"], "新兴市场"),
    (["黄金", "贵金属", "白银"], "黄金/贵金属"),
    (["原油", "石油", "天然气", "能源", "商品", "资源"], "能源/商品"),
    (["REITs", "REIT", "房地产"], "海外REITs"),
    (["债券", "债"], "海外债券"),
    (["医疗", "医药", "生物", "健康"], "全球医药"),
    (["科技", "半导体", "芯片", "人工智能", "互联网", "互联", "消费", "新能源", "高端制造", "信息", "数字", "军工", "汽车"], "全球行业主题"),
    (["亚太", "亚洲"], "亚太"),
    (["全球", "世界", "海外", "环球", "国际", "中华"], "全球"),
]


def classify_qdii(ftype):
    """将天天基金 FTYPE 映射到5个 QDII 大类"""
    f = (ftype or "").strip()
    if f in ("指数型-海外股票", "指数型-股票"):
        return "QDII-被动权益"
    if f in ("QDII-普通股票", "QDII-混合偏股", "QDII-混合灵活", "QDII-混合平衡", "QDII-FOF"):
        return "QDII-主动权益"
    if f in ("QDII-纯债", "QDII-混合债"):
        return "QDII纯债"
    if f in ("QDII-商品", "QDII-REITs"):
        return "QDII商品"
    return "其他"


def infer_region(name, ftype):
    t = (name or "") + " " + (ftype or "")
    for kws, region in REGION_RULES:
        for k in kws:
            if k in t:
                return region
    return "未细分（主动配置）"


# ---------- 额度解析 ----------
def parse_limit(sgzt, minsg=None):
    """返回 (额度数值 or None, 额度显示文本, 申购状态归类)"""
    s = (sgzt or "").strip()
    # 接口未返回申购状态（多为天天基金不代销的份额）
    if not s:
        return None, "未披露", "未披露"

    # 申购状态 = 括号前的主状态
    prefix = re.split(r"[（(]", s)[0].strip()
    if prefix not in ("开放申购", "限大额", "暂停申购", "封闭期", "场内交易"):
        if "封闭" in s:
            prefix = "封闭期"
        elif "场内交易" in s:
            prefix = "场内交易"
        elif "暂停申购" in s:
            prefix = "暂停申购"
        elif "限大额" in s:
            prefix = "限大额"
        elif "开放申购" in s:
            prefix = "开放申购"
        else:
            prefix = s or "未知"

    # 提取“单日投资上限X元” / “单日累计购买上限X元”
    m = re.search(r"单日(?:投资上限|累计购买上限|申购上限)\s*([\d.]+)\s*(万)?\s*元", s)
    limit_val = None
    limit_txt = None
    if m:
        v = float(m.group(1))
        if m.group(2):
            v *= 10000
        limit_val = v
        limit_txt = ("%.0f" % v) if v >= 1 else ("%.2f" % v)

    if limit_val is None:
        if prefix == "开放申购":
            limit_txt = "不限"
        elif prefix == "暂停申购":
            limit_txt = "暂停申购"
        elif prefix == "封闭期":
            limit_txt = "封闭期"
        elif prefix == "限大额":
            limit_txt = "限大额（未披露额度）"
        else:
            limit_txt = prefix

    # 暂停申购的，额度无实际意义
    if prefix == "暂停申购":
        limit_val = None
        limit_txt = "暂停申购"

    return limit_val, limit_txt, prefix


def fmt_size(v):
    try:
        f = float(v)
        if f <= 0:
            return None
        return round(f / 1e8, 2)
    except Exception:
        return None


def main():
    data = json.load(open(SRC, encoding="utf-8"))
    rows = []
    for d in data:
        code = d.get("FCODE")
        name = (d.get("SHORTNAME") or "").strip()
        if not name:
            continue
        sgzt = d.get("SGZT")
        minsg = d.get("MINSG")
        lv, ltxt, status = parse_limit(sgzt, minsg)
        ftype = d.get("FTYPE") or ""
        def num(k):
            v = d.get(k)
            try:
                f = float(v)
                return round(f, 2)
            except Exception:
                return None
        rows.append({
            "code": code,
            "name": name,
            "ftype": ftype,
            "region": infer_region(name, ftype),
            "limit_val": lv,
            "limit_txt": ltxt,
            "status": status,
            "sgzt_raw": sgzt or "",
            "ret_1n": num("SYL_1N"),
            "ret_jn": None,
            "ret_1y_m": num("SYL_Y"),
            "ret_1d": num("RZDF"),
            "nav": d.get("DWJZ") or "",
            "nav_date": d.get("FSRQ") or "",
            "company": d.get("JJGS") or "",
            "size": fmt_size(d.get("ENDNAV")),
            "manager": d.get("JJJL") or "",
        })

    # 排序：额度从高到低（不限 > 大额 > 小额 > 封闭 > 暂停 > 场内）
    rank_of = {"开放申购": 0, "限大额": 1, "封闭期": 2, "暂停申购": 3, "场内交易": 4}

    def sort_key(r):
        g = rank_of.get(r["status"], 5)
        if g == 1 and r["limit_val"] is not None:
            return (g, -r["limit_val"], r["name"])
        return (g, 0, r["name"])

    rows.sort(key=sort_key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "额度排行"

    headers = ["排名", "基金代码", "基金名称", "类型", "主要投资区域\n（按名称推断）",
               "单日申购上限(元)", "额度显示", "申购状态",
               "日涨跌幅(%)", "近1月(%)", "近1年收益率(%)",
               "单位净值", "净值日期", "基金公司", "基金规模(亿元)", "申购状态原文"]
    ws.append(headers)

    hdr_fill = PatternFill("solid", fgColor="1EB9E1")
    hdr_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10)
    thin = Side(style="thin", color="D3D3D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 34

    fills = {
        "开放申购": PatternFill("solid", fgColor="E8F7EC"),
        "限大额": PatternFill("solid", fgColor="FFF6E5"),
        "暂停申购": PatternFill("solid", fgColor="FBE9E7"),
    }

    for i, r in enumerate(rows, 1):
        ws.append([
            i, r["code"], r["name"], r["ftype"], r["region"],
            (r["limit_val"] if r["limit_val"] is not None else ""),
            r["limit_txt"], r["status"],
            (r["ret_1d"] if r["ret_1d"] is not None else ""),
            (r["ret_1y_m"] if r["ret_1y_m"] is not None else ""),
            (r["ret_1n"] if r["ret_1n"] is not None else ""),
            r["nav"], r["nav_date"], r["company"],
            (r["size"] if r["size"] is not None else ""),
            r["sgzt_raw"],
        ])
        row_i = i + 1
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_i, column=c)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if c == 3 else "center", vertical="center")
        f = fills.get(r["status"])
        if f:
            for c in range(1, len(headers) + 1):
                ws.cell(row=row_i, column=c).fill = f

    widths = [6, 10, 34, 16, 18, 16, 12, 12, 11, 10, 13, 10, 12, 14, 13, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # ---- 汇总 sheet ----
    ws2 = wb.create_sheet("汇总统计")
    c1 = Counter(r["status"] for r in rows)
    c2 = Counter(r["region"] for r in rows)
    ws2.append(["申购状态分布", "数量"])
    for k, v in c1.most_common():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["主要投资区域分布", "数量"])
    for k, v in c2.most_common():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["数据来源", "天天基金移动端接口 fundmobapi.eastmoney.com"])
    ws2.append(["取数日期", time.strftime("%Y-%m-%d")])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=2):
        for cell in row:
            cell.font = body_font
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 42

    wb.save(DST)
    print("saved:", DST)
    print("rows:", len(rows))
    print("status:", c1.most_common())
    print("region:", c2.most_common())
    # 预览
    for r in rows[:12]:
        print(r["code"], r["name"][:26], "|", r["region"], "|", r["limit_txt"], "|", r["status"])


if __name__ == "__main__":
    main()
